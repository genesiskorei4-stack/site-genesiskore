"""Gera a planilha modelo de DRE gerencial para distribuidora, que a página
/dre-gerencial-para-distribuidora/ oferece para baixar.

Uso: python3 execution/gerar_modelo_dre.py
Saída: dre-gerencial-para-distribuidora/modelo-dre-gerencial-distribuidora.xlsx

Três abas: "DRE gerencial" (12 meses, fórmulas prontas), "Margem por produto" e "Como usar".
Os valores de saída entram positivos; as fórmulas subtraem. Nenhum número de cliente: a planilha
vai em branco. Feita em 30/10/2026, no plano de SEO e GEO (pedido do Felipe).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / 'dre-gerencial-para-distribuidora' / 'modelo-dre-gerencial-distribuidora.xlsx'

MESES = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
PRETO = '000000'
CINZA = 'F0EFEC'
FINO = Side(style='thin', color='D4D4D4')

# (rótulo, tipo, chave). tipo: 'entra', 'sai', 'resultado' (fórmula), 'percentual' (fórmula), 'titulo'.
LINHAS = [
    ('Receita', 'titulo', None),
    ('Receita bruta de vendas', 'entra', 'bruta'),
    ('Devoluções', 'sai', 'devolucoes'),
    ('Impostos sobre a venda (ICMS, PIS, Cofins, ou o DAS no Simples)', 'sai', 'impostos'),
    ('Receita líquida', 'resultado', 'liquida'),
    ('Mercadoria', 'titulo', None),
    ('Custo da mercadoria vendida (CMV)', 'sai', 'cmv'),
    ('Lucro bruto', 'resultado', 'bruto'),
    ('Margem bruta (% da receita líquida)', 'percentual', 'p_bruto'),
    ('Despesas que andam com a venda', 'titulo', None),
    ('Frete de entrega', 'sai', 'frete'),
    ('Comissão de vendedores e representantes', 'sai', 'comissao'),
    ('Bonificação, troca e avaria', 'sai', 'bonificacao'),
    ('Taxas de cartão, boleto e marketplace', 'sai', 'taxas'),
    ('Margem de contribuição', 'resultado', 'contribuicao'),
    ('Margem de contribuição (% da receita líquida)', 'percentual', 'p_contribuicao'),
    ('Despesas fixas', 'titulo', None),
    ('Pessoal (salários, encargos e pró-labore)', 'sai', 'pessoal'),
    ('Galpão (aluguel, energia, segurança)', 'sai', 'galpao'),
    ('Frota própria (combustível, manutenção, seguro)', 'sai', 'frota'),
    ('Sistemas, contador e serviços', 'sai', 'servicos'),
    ('Outras despesas fixas', 'sai', 'outras'),
    ('Resultado operacional', 'resultado', 'operacional'),
    ('Resultado operacional (% da receita líquida)', 'percentual', 'p_operacional'),
    ('Financeiro e impostos sobre o lucro', 'titulo', None),
    ('Juros e antecipação de recebíveis', 'sai', 'juros'),
    ('Receitas financeiras', 'entra', 'rec_fin'),
    ('IRPJ e CSLL (zero no Simples, que já está no DAS)', 'sai', 'irpj'),
    ('Resultado líquido', 'resultado', 'liquido'),
    ('Margem líquida (% da receita líquida)', 'percentual', 'p_liquido'),
]

# Como cada resultado e cada percentual se calcula, pelas chaves das linhas acima.
CONTAS = {
    'liquida': ('bruta', ['devolucoes', 'impostos'], []),
    'bruto': ('liquida', ['cmv'], []),
    'contribuicao': ('bruto', ['frete', 'comissao', 'bonificacao', 'taxas'], []),
    'operacional': ('contribuicao', ['pessoal', 'galpao', 'frota', 'servicos', 'outras'], []),
    'liquido': ('operacional', ['juros', 'irpj'], ['rec_fin']),
}
PERCENTUAIS = {'p_bruto': 'bruto', 'p_contribuicao': 'contribuicao', 'p_operacional': 'operacional', 'p_liquido': 'liquido'}


def aba_dre(wb):
    ws = wb.active
    ws.title = 'DRE gerencial'
    ws['A1'] = 'DRE gerencial da distribuidora'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Preencha os valores de cada mês em reais, todos positivos. As linhas em cinza se calculam sozinhas.'
    ws['A2'].font = Font(italic=True, color='555555')

    cab = 4
    ws.cell(row=cab, column=1, value='Linha')
    ws.cell(row=cab, column=2, value='Tipo')
    for i, mes in enumerate(MESES):
        ws.cell(row=cab, column=3 + i, value=mes)
    ws.cell(row=cab, column=15, value='Ano')
    for c in range(1, 16):
        cel = ws.cell(row=cab, column=c)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor=PRETO)
        cel.alignment = Alignment(horizontal='center' if c > 2 else 'left')

    linha_de = {}
    r = cab + 1
    for rotulo, tipo, chave in LINHAS:
        if chave:
            linha_de[chave] = r
        r += 1

    r = cab + 1
    for rotulo, tipo, chave in LINHAS:
        ws.cell(row=r, column=1, value=rotulo)
        if tipo == 'titulo':
            ws.cell(row=r, column=1).font = Font(bold=True)
            r += 1
            continue
        ws.cell(row=r, column=2, value={'entra': 'Entra', 'sai': 'Sai', 'resultado': 'Resultado', 'percentual': '%'}[tipo])
        for c in range(3, 16):
            col = get_column_letter(c)
            cel = ws.cell(row=r, column=c)
            if tipo in ('entra', 'sai'):
                if c == 15:
                    cel.value = f'=SUM(C{r}:N{r})'
                cel.number_format = '#,##0.00'
            elif tipo == 'resultado':
                base, menos, mais = CONTAS[chave]
                formula = f'={col}{linha_de[base]}'
                formula += ''.join(f'-{col}{linha_de[k]}' for k in menos)
                formula += ''.join(f'+{col}{linha_de[k]}' for k in mais)
                cel.value = formula
                cel.number_format = '#,##0.00'
                cel.font = Font(bold=True)
                cel.fill = PatternFill('solid', fgColor=CINZA)
            else:
                alvo = linha_de[PERCENTUAIS[chave]]
                liq = linha_de['liquida']
                cel.value = f'=IF({col}{liq}=0,"",{col}{alvo}/{col}{liq})'
                cel.number_format = '0.0%'
                cel.fill = PatternFill('solid', fgColor=CINZA)
            cel.border = Border(bottom=FINO)
        if tipo in ('resultado', 'percentual'):
            ws.cell(row=r, column=1).font = Font(bold=(tipo == 'resultado'))
            ws.cell(row=r, column=1).fill = PatternFill('solid', fgColor=CINZA)
            ws.cell(row=r, column=2).fill = PatternFill('solid', fgColor=CINZA)
        r += 1

    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 11
    for c in range(3, 16):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = ws.cell(row=cab + 1, column=3)


def aba_produto(wb):
    ws = wb.create_sheet('Margem por produto')
    ws['A1'] = 'Margem de contribuição por produto, num mês'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Uma linha por produto. Preço, custo, frete e comissão por unidade; imposto em % do preço. As colunas em cinza se calculam sozinhas.'
    ws['A2'].font = Font(italic=True, color='555555')
    cab = ['Produto', 'Quantidade vendida', 'Preço médio por unidade', 'Custo por unidade', 'Imposto sobre a venda (%)',
           'Frete por unidade', 'Comissão (%)', 'Bonificação e troca por unidade', 'Margem por unidade',
           'Margem (% do preço)', 'Margem no mês']
    for i, t in enumerate(cab, start=1):
        cel = ws.cell(row=4, column=i, value=t)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor=PRETO)
        cel.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[4].height = 46
    for r in range(5, 55):
        ws.cell(row=r, column=9, value=f'=IF(C{r}="","",C{r}-D{r}-C{r}*E{r}-F{r}-C{r}*G{r}-H{r})')
        ws.cell(row=r, column=10, value=f'=IF(OR(C{r}="",C{r}=0),"",I{r}/C{r})')
        ws.cell(row=r, column=11, value=f'=IF(I{r}="","",I{r}*B{r})')
        for c in (3, 4, 6, 8, 9, 11):
            ws.cell(row=r, column=c).number_format = '#,##0.00'
        for c in (5, 7, 10):
            ws.cell(row=r, column=c).number_format = '0.0%'
        for c in (9, 10, 11):
            ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=CINZA)
    ws.column_dimensions['A'].width = 34
    for c in range(2, 12):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.freeze_panes = 'B5'


def aba_como_usar(wb):
    ws = wb.create_sheet('Como usar')
    textos = [
        ('Como usar este modelo', True),
        ('', False),
        ('1. Na aba DRE gerencial, preencha cada mês com os valores em reais, todos positivos. A planilha subtrai o que sai.', False),
        ('2. A receita, as devoluções e o imposto vêm do faturamento do ERP ou do relatório do contador.', False),
        ('3. O custo da mercadoria vendida é o custo do que saiu no mês, e não o que foi comprado no mês.', False),
        ('4. Frete, comissão, bonificação e taxas ficam separados das despesas fixas, porque crescem junto com a venda.', False),
        ('5. Na aba Margem por produto, uma linha por produto mostra quanto cada um deixa por unidade e no mês.', False),
        ('6. Compare os meses lado a lado: a pergunta não é quanto vendeu, é quanto sobrou e por quê.', False),
        ('', False),
        ('Regime tributário e regra de cada produto são com o seu contador. Este modelo é gerencial: serve para decidir, e não substitui o DRE contábil.', False),
        ('A partir de janeiro de 2027 a CBS passa a ser cobrada no lugar do PIS e da Cofins (Lei Complementar 214/2025). A linha de impostos sobre a venda muda de conta: confirme com o contador.', False),
        ('', False),
        ('Modelo da GenesisKore, de São Paulo. Sistemas sob medida e consultoria de gestão: https://genesiskore.com/dre-gerencial-para-distribuidora/', False),
    ]
    for i, (t, negrito) in enumerate(textos, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=negrito, size=14 if negrito else 11)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 120


def main():
    wb = Workbook()
    aba_dre(wb)
    aba_produto(wb)
    aba_como_usar(wb)
    wb.properties.creator = 'GenesisKore'
    wb.properties.title = 'Modelo de DRE gerencial para distribuidora'
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print('Gravado:', SAIDA.relative_to(RAIZ))


if __name__ == '__main__':
    main()
